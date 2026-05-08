from openpyxl import load_workbook

output_file = r"F:\PROJECTS\lepure-business-opportunity\analysis-ollama\Jan_Mar_2026_CGT_Opportunities_Formatted.xlsx"

try:
    wb = load_workbook(output_file)
    ws = wb.active
    print("Row 2, Column 9 value (新闻原文):")
    print(ws.cell(row=2, column=9).value)
    print("Row 3, Column 9 value (新闻原文):")
    print(ws.cell(row=3, column=9).value)
except Exception as e:
    print(f"Error: {e}")
