import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import os

source_file = r"F:\PROJECTS\lepure-business-opportunity\analysis-ollama\Jan_Mar_2026_CGT_Opportunities.xlsx"
output_file = r"F:\PROJECTS\lepure-business-opportunity\analysis-ollama\Jan_Mar_2026_CGT_Opportunities_Formatted.xlsx"

# 1. Load data
df = pd.read_excel(source_file)

# 2. Data Cleaning: Remove rows where '公司名称' is empty or NaN
df = df[df['公司名称'].notna() & (df['公司名称'].astype(str).str.strip() != '')]

# 3. Prepare Formatted Data
formatted_data = []

for _, row in df.iterrows():
    # Handle NaN values for strings
    company = str(row['公司名称']) if pd.notna(row['公司名称']) else ""
    province = str(row['省份']) if pd.notna(row['省份']) else "--"
    city = str(row['城市']) if pd.notna(row['城市']) else "--"
    client_type = str(row['客户类型']) if pd.notna(row['客户类型']) else "--"
    app_scene = str(row['应用场景']) if pd.notna(row['应用场景']) else "--"
    project_name = str(row['项目名称']) if pd.notna(row['项目名称']) else "--"
    project_stage = str(row['项目阶段']) if pd.notna(row['项目阶段']) else "--"
    
    # Newspaper/Link handling - using Excel HYPERLINK formula
    title = str(row['文章标题']) if pd.notna(row['文章标题']) else "链接"
    link = str(row['文章链接']) if pd.notna(row['文章链接']) else ""
    news_content = f'=HYPERLINK("{link}", "{title}")' if link else title
    
    pub_time = str(row['发布时间']) if pd.notna(row['发布时间']) else ""
    
    formatted_data.append({
        "相关企业": company,
        "所在省份": province,
        "所在城市": city,
        "客户类型": client_type,
        "应用场景": app_scene,
        "项目管线": project_name,
        "临床阶段": project_stage,
        "价值摘要": "--", # Added empty column to match image
        "新闻原文": news_content,
        "发布时间": pub_time
    })

df_new = pd.DataFrame(formatted_data)

# 4. Create Workbook and apply styling
wb = Workbook()
ws = wb.active
ws.title = "CGT Opportunities"

# Headers
headers = list(df_new.columns)
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)

# Data
for r_idx, row_dict in enumerate(formatted_data, 2):
    for c_idx, (key, value) in enumerate(row_dict.items(), 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        if key == "新闻原文" and value.startswith("=HYPERLINK"):
            cell.font = Font(color="0000FF", underline="single")

# Styling
header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # Light green/greyish
header_font = Font(bold=True, size=11)
centered_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
border_side = Side(border_style="thin", color="000000")
full_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

# Apply to headers
for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = centered_alignment
    cell.border = full_border

# Apply to data
for row in range(2, len(formatted_data) + 2):
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = full_border
        if headers[col-1] in ["相关企业", "所在省份", "所在城市", "客户类型", "临床阶段", "发布时间"]:
            cell.alignment = centered_alignment
        else:
            cell.alignment = left_alignment

# Adjust column widths
column_widths = {
    "相关企业": 30,
    "所在省份": 15,
    "所在城市": 15,
    "客户类型": 15,
    "应用场景": 25,
    "项目管线": 25,
    "临床阶段": 15,
    "价值摘要": 40,
    "新闻原文": 40,
    "发布时间": 20
}

for col_num, header in enumerate(headers, 1):
    ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = column_widths.get(header, 20)

# Save
wb.save(output_file)
print(f"Formatted file saved to: {output_file}")
